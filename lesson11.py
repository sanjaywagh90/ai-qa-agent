# Lesson 11 - Export Test Case Report to Excel (with all field types)

import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment


# ---------- SQL TOOL ----------
def fetch_requirements_from_db():
    conn = sqlite3.connect("testing.db")
    cursor = conn.cursor()
    cursor.execute("SELECT field_name, field_type, min_len, max_len FROM requirements")
    rows = cursor.fetchall()
    conn.close()

    requirements = []
    for row in rows:
        requirements.append({
            "field_name": row[0],
            "field_type": row[1],
            "min_len": row[2],
            "max_len": row[3]
        })
    return requirements


# ---------- TOOL FOR TEXT FIELD ----------
def generate_text_field_cases(min_len, max_len):
    positive = [
        f"Enter {min_len} alphabets (minimum boundary)",
        f"Enter {max_len} alphabets (maximum boundary)",
        "Enter valid alphabets within range"
    ]
    negative = [
        "Leave field blank",
        f"Enter {min_len - 1} alphabets (below minimum)",
        f"Enter {max_len + 1} alphabets (above maximum)",
        "Enter numbers instead of alphabets",
        "Enter special characters"
    ]
    return positive, negative


# ---------- TOOL FOR EMAIL FIELD ----------
def generate_email_field_cases(min_len=None, max_len=None):
    positive = [
        "Enter valid email (e.g. user@example.com)",
        "Enter email with subdomain (e.g. user@mail.example.com)"
    ]
    negative = [
        "Leave field blank",
        "Enter email without @ symbol",
        "Enter email without domain (e.g. user@)",
        "Enter email with spaces",
        "Enter email with special characters not allowed"
    ]
    return positive, negative


# ---------- TOOL FOR PASSWORD FIELD ----------
def generate_password_field_cases(min_len, max_len):
    positive = [
        f"Enter valid password with {min_len}-{max_len} chars, uppercase, number, special char",
        "Enter password exactly at minimum length with all required character types"
    ]
    negative = [
        "Leave field blank",
        f"Enter password below {min_len} characters",
        f"Enter password above {max_len} characters",
        "Enter password without uppercase letter",
        "Enter password without number",
        "Enter password without special character"
    ]
    return positive, negative


# ---------- TOOL FOR NUMBER FIELD ----------
def generate_number_field_cases(min_len, max_len):
    positive = [
        f"Enter minimum valid number ({min_len})",
        f"Enter maximum valid number ({max_len})",
        "Enter a valid number within range"
    ]
    negative = [
        "Leave field blank",
        f"Enter number below minimum ({min_len - 1})",
        f"Enter number above maximum ({max_len + 1})",
        "Enter alphabets instead of numbers",
        "Enter special characters instead of numbers",
        "Enter decimal value if only integers are allowed"
    ]
    return positive, negative


# ---------- TOOL FOR DATE FIELD ----------
def generate_date_field_cases(min_len=None, max_len=None):
    positive = [
        "Enter valid date in correct format (e.g. DD/MM/YYYY)",
        "Enter today's date",
        "Enter a valid past date (if allowed)"
    ]
    negative = [
        "Leave field blank",
        "Enter date in wrong format (e.g. MM-DD-YYYY when DD/MM/YYYY expected)",
        "Enter invalid date (e.g. 32/13/2025)",
        "Enter future date (if not allowed)",
        "Enter alphabets instead of a date",
        "Enter date with special characters"
    ]
    return positive, negative


# ---------- TOOL FOR DROPDOWN FIELD ----------
def generate_dropdown_field_cases(min_len=None, max_len=None):
    positive = [
        "Select a valid option from the dropdown",
        "Select the first option in the list",
        "Select the last option in the list"
    ]
    negative = [
        "Submit form without selecting any option (if mandatory)",
        "Verify default placeholder text is not a valid submission",
        "Verify dropdown list is not editable manually"
    ]
    return positive, negative


# ---------- TOOL REGISTRY ----------
field_tools = {
    "text": generate_text_field_cases,
    "email": generate_email_field_cases,
    "password": generate_password_field_cases,
    "number": generate_number_field_cases,
    "date": generate_date_field_cases,
    "dropdown": generate_dropdown_field_cases
}


# ---------- AGENT ----------
class QAAgent:
    def __init__(self, field_tools):
        self.field_tools = field_tools

    def run(self):
        requirements = fetch_requirements_from_db()
        final_report = []

        for req in requirements:
            field_name = req["field_name"]
            field_type = req["field_type"]
            min_len = req["min_len"]
            max_len = req["max_len"]

            tool = self.field_tools.get(field_type)
            if not tool:
                print(f"No tool found for field type: {field_type} (field: {field_name})")
                continue

            positive, negative = tool(min_len, max_len)

            final_report.append({
                "field_name": field_name,
                "positive": positive,
                "negative": negative
            })

        return final_report


# ---------- EXCEL EXPORT TOOL ----------
def export_report_to_excel(report, filename="Test_Case_Report.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    headers = ["Sr No", "Field Name", "Test Case Type", "Test Case Description"]
    ws.append(headers)

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    sr_no = 1
    for item in report:
        for case in item["positive"]:
            ws.append([sr_no, item["field_name"], "Positive", case])
            sr_no += 1
        for case in item["negative"]:
            ws.append([sr_no, item["field_name"], "Negative", case])
            sr_no += 1

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 65

    wb.save(filename)
    print(f"Report saved successfully as '{filename}'")


# ---------- RUN ----------
print("Agent Thinking: Generating test cases and exporting to Excel...\n")

agent = QAAgent(field_tools)
report = agent.run()

export_report_to_excel(report)